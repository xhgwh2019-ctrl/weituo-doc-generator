"""核心生成逻辑：占位符替换 + 多文档合并为同一个文件。"""
import copy
import datetime
import re
import sys
from pathlib import Path

from paths import templates_dir, output_dir

TEMPLATES_DIR = templates_dir()
OUTPUT_DIR = output_dir()

# 模板顺序（合并后的章节顺序）
ORDER = [
    ("01_聘请律师合同.docx", "法律事务聘请律师合同"),
    ("02_接待笔录.docx", "接待笔录"),
    ("03_风险告知书.docx", "诉讼风险告知书"),
    ("04_授权委托书.docx", "授权委托书"),
    ("05_账户信息确认书.docx", "账户信息确认书"),
    ("06_送达地址确认书.docx", "送达地址确认书"),
]

CN_DIGITS = "零壹贰叁肆伍陆柒捌玖"
CN_UNITS = ["", "拾", "佰", "仟"]
CN_BIG = ["", "万", "亿", "兆"]


def amount_to_cn(n) -> str:
    """6000 -> 陆仟元整；支持小数（角分）"""
    try:
        n = float(str(n).replace(",", "").replace("￥", "").strip())
    except Exception:
        return ""
    if n == 0:
        return "零元整"
    integer = int(n)
    jiao_fen = round((n - integer) * 100)
    # 整数部分
    s = str(integer)
    res = ""
    # 按4位分组
    groups = []
    while s:
        groups.append(s[-4:])
        s = s[:-4]
    for gi, g in enumerate(groups):
        g = g.zfill(4)
        tmp = ""
        for i, ch in enumerate(g):
            d = int(ch)
            u = CN_UNITS[4 - 1 - i]
            if d == 0:
                if tmp and not tmp.endswith("零"):
                    tmp += "零"
            else:
                tmp += CN_DIGITS[d] + u
        tmp = tmp.rstrip("零")
        if tmp:
            tmp += CN_BIG[gi]
        res = tmp + res
    res = re.sub(r"零+", "零", res).strip("零")
    if not res:
        res = "零"
    res += "元"
    if jiao_fen == 0:
        res += "整"
    else:
        j = jiao_fen // 10
        f = jiao_fen % 10
        if j:
            res += CN_DIGITS[j] + "角"
        if f:
            res += CN_DIGITS[f] + "分"
    return res


def normalize_case_reason(s: str) -> str:
    """统一案由：去掉首尾空格和'纠纷'后缀，返回不带'纠纷'的案由。生成时再加'纠纷'。"""
    s = (s or "").strip()
    s = re.sub(r"[，,。.；;]+$", "", s)
    if s.endswith("纠纷"):
        s = s[:-2]
    return s


def build_mapping(data: dict, cfg: dict) -> dict:
    today = data.get("date") or datetime.date.today()
    if isinstance(today, str):
        today = datetime.date.fromisoformat(today)
    y, m, d = today.year, today.month, today.day
    date_str = f"{y}年{m}月{d}日"

    client_names = (data.get("client_names") or "").strip()
    opponent = (data.get("opponent") or "").strip()
    reason = normalize_case_reason(data.get("case_reason") or "")
    reason_full = f"{reason}纠纷" if reason else ""
    # 与对方+案由 的完整表述，如：与王艳萍土地承包经营权转让合同纠纷
    vs_text = f"与{opponent}{reason_full}" if opponent else reason_full

    fee = str(data.get("fee") or "").strip()
    fee_cn = amount_to_cn(fee) if fee else ""
    other_fee = str(data.get("other_fee") or "0").strip()

    contract_no = (data.get("contract_no") or "").strip()
    if not contract_no:
        prefix = cfg.get("contract_no_prefix", "（）冀禄律民字第")
        contract_no = f"{prefix}    号"

    case_no = (data.get("case_no") or "").strip()
    court = (data.get("court") or cfg.get("court_default", "")).strip()

    bank_holder = (data.get("bank_holder") or client_names).strip()
    valid_until = (data.get("valid_until") or cfg.get("delegate_period_default", "一审终结之日止")).strip()

    m_map = {
        "{{client_names}}": client_names,
        "{{client_ids}}": (data.get("client_ids") or "").strip(),
        "{{client_addr}}": (data.get("client_addr") or "").strip(),
        "{{client_phone}}": (data.get("client_phone") or "").strip(),
        "{{opponent}}": opponent,
        "{{case_reason}}": reason,
        "{{case_reason_full}}": reason_full,
        "{{vs_text}}": vs_text,
        "{{contract_no}}": contract_no,
        "{{fee}}": fee,
        "{{fee_cn}}": fee_cn,
        "{{other_fee}}": other_fee,
        "{{fee_method}}": data.get("fee_method") or cfg.get("fee_method", "现金"),
        "{{pay_method}}": data.get("pay_method") or cfg.get("pay_method", "签订合同之日支付"),
        "{{fee_standard}}": cfg.get("fee_standard", ""),
        "{{bank_holder}}": bank_holder,
        "{{bank_account}}": (data.get("bank_account") or "").strip(),
        "{{bank_branch}}": (data.get("bank_branch") or "").strip(),
        "{{case_no}}": case_no,
        "{{court}}": court,
        "{{valid_until}}": valid_until,
        "{{date_str}}": date_str,
        "{{date_y}}": str(y),
        "{{date_m}}": str(m),
        "{{date_d}}": str(d),
        "{{law_firm}}": cfg.get("law_firm", ""),
        "{{lawyer}}": cfg.get("lawyer", ""),
        "{{lawyer_id}}": cfg.get("lawyer_id", ""),
        "{{lawyer_phone}}": cfg.get("lawyer_phone", ""),
        "{{lawyer_addr}}": cfg.get("lawyer_addr", ""),
        "{{office_addr}}": cfg.get("office_addr", ""),
    }
    return m_map


def replace_in_paragraph(paragraph, mapping: dict) -> bool:
    """占位符替换（保持 run 边界与格式：下划线/加粗/字体不串位）。"""
    runs = paragraph.runs
    if not runs:
        return False
    changed = False
    for old, new in mapping.items():
        if not old:
            continue
        while True:
            full = "".join(r.text for r in runs)
            idx = full.find(old)
            if idx < 0:
                break
            end = idx + len(old)
            pos = 0
            ri = rj = None
            off_s = off_e = 0
            for i, r in enumerate(runs):
                L = len(r.text)
                if ri is None and idx < pos + L:
                    ri, off_s = i, idx - pos
                if end - 1 < pos + L:
                    rj, off_e = i, end - pos
                    break
                pos += L
            if ri is None or rj is None:
                break
            if ri == rj:
                runs[ri].text = runs[ri].text[:off_s] + new + runs[ri].text[off_e:]
            else:
                runs[ri].text = runs[ri].text[:off_s] + new
                for k in range(ri + 1, rj):
                    runs[k].text = ""
                runs[rj].text = runs[rj].text[off_e:]
            changed = True
    return changed


def replace_in_doc(doc, mapping: dict):
    for p in doc.paragraphs:
        replace_in_paragraph(p, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, mapping)
    # 页眉页脚
    for section in doc.sections:
        if section.header:
            for p in section.header.paragraphs:
                replace_in_paragraph(p, mapping)
        if section.footer:
            for p in section.footer.paragraphs:
                replace_in_paragraph(p, mapping)


def _is_empty_para(p) -> bool:
    return not (p.text or "").strip()


def _has_sect_break(p) -> bool:
    """该段落是否内含分节符（合一版的分页标记段），绝不可删除。"""
    from docx.oxml.ns import qn
    pPr = p._p.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:sectPr')) is not None


def strip_doc(doc, keep_one_blank=False):
    """去掉顶层首尾空段落；连续空段落最多保留1个（默认全压掉连续空白）。
    内含分节符的分页标记段永远保留。"""
    # 去头部空段（至少保留1段，避免 body 全空）
    while (len(doc.paragraphs) > 1 and _is_empty_para(doc.paragraphs[0])
           and not _has_sect_break(doc.paragraphs[0])):
        doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p)
    # 去尾部空段
    while (len(doc.paragraphs) > 1 and _is_empty_para(doc.paragraphs[-1])
           and not _has_sect_break(doc.paragraphs[-1])):
        doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
    # 压缩中间连续空段：最多保留1个
    prev_empty = False
    for p in list(doc.paragraphs):
        if _has_sect_break(p):
            prev_empty = False
            continue
        if _is_empty_para(p):
            if prev_empty and len(doc.paragraphs) > 1:
                p._p.getparent().remove(p._p)
            else:
                prev_empty = True
        else:
            prev_empty = False
    return doc


def _new_clean_doc():
    """以 02 接待笔录模板为底新建空文档并清空 body。
    02/03/05 的文字没有显式字体、依赖 WPS 主题与默认样式；用 python-docx
    空文档做底会导致字体混乱、分页错乱（如笔录 1 页变 2 页）。"""
    from docx import Document
    doc = Document(str(TEMPLATES_DIR / ORDER[1][0]))
    body = doc.element.body
    for el in list(body):
        body.remove(el)
    return doc


def _bare_paragraph(merged):
    """新建一个不带样式引用的空段落（02 底包没有 Normal 样式，add_paragraph 会 KeyError）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    p_el = OxmlElement('w:p')
    body = merged.element.body
    sectPr = body.find(qn('w:sectPr'))
    if sectPr is not None:
        sectPr.addprevious(p_el)
    else:
        body.append(p_el)
    return Paragraph(p_el, merged)


def _section_break(merged, src_doc):
    """在 merged 末尾追加一个“奇数页起新节”分节符，并套用 src_doc 的版式。
    效果：下一份文档强制新开一页；双面打印时每份从正面（奇数页）开始，
    若上一份恰好占满奇数页，Word 会自动垫一张空白偶数页。"""
    import copy as _copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = _bare_paragraph(merged)
    pPr = p._p.get_or_add_pPr()
    sectPr = OxmlElement('w:sectPr')
    t = OxmlElement('w:type')
    t.set(qn('w:val'), 'oddPage')
    sectPr.append(t)
    src_sect = src_doc.sections[0]._sectPr if src_doc.sections else None
    if src_sect is not None:
        for tag in ('w:pgSz', 'w:pgMar', 'w:cols', 'w:docGrid'):
            el = src_sect.find(qn(tag))
            if el is not None:
                sectPr.append(_copy.deepcopy(el))
    pPr.append(sectPr)
    return p


def _clear_footers(doc):
    """清空全部三种页脚（默认/首页/偶数页，含页码域）。用户要求不显示页码；模板文件本身不动。"""
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            try:
                for p in footer.paragraphs:
                    for r in p.runs:
                        r.text = ""
            except Exception:
                pass
    return doc


def _apply_section_to_body(merged, src_doc):
    """把 src_doc 的版式设为合一版 body 级（最后一节）的版式。
    body 级若无 sectPr（新底已清空）则新建一个。"""
    import copy as _copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    body = merged.element.body
    body_sect = body.find(qn('w:sectPr'))
    if body_sect is None:
        body_sect = OxmlElement('w:sectPr')
        body.append(body_sect)
    src_sect = src_doc.sections[0]._sectPr if src_doc.sections else None
    if src_sect is None:
        return
    for tag in ('w:pgSz', 'w:pgMar', 'w:cols', 'w:docGrid'):
        old = body_sect.find(qn(tag))
        if old is not None:
            body_sect.remove(old)
        el = src_sect.find(qn(tag))
        if el is not None:
            body_sect.append(_copy.deepcopy(el))


def _append_doc(merged, src_doc):
    """把已填充好的 src_doc 的 body 元素原样拼接到 merged（版式由分节符承载）。
    注意：必须插到 body 级 sectPr 之前，直接 append 会跑到 sectPr 后面导致错乱。"""
    from docx.oxml.ns import qn
    body = merged.element.body
    sectPr = body.find(qn('w:sectPr'))
    for el in list(src_doc.element.body):
        if el.tag.endswith("sectPr"):
            continue
        c = copy.deepcopy(el)
        if sectPr is not None:
            sectPr.addprevious(c)
        else:
            body.append(c)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _clone_contract_hf(merged_path: Path):
    """把 01 合同模板的 logo 页眉克隆进合一版第 1 节；其余节挂空白页眉。
    页脚一律不挂（用户要求不显示页码）。只动 zip 结构，不动正文。"""
    import zipfile
    import xml.etree.ElementTree as ET
    ET.register_namespace('w', W_NS)
    ET.register_namespace('r', R_NS)
    src_tpl = TEMPLATES_DIR / ORDER[0][0]
    if not src_tpl.exists():
        return
    with zipfile.ZipFile(str(src_tpl)) as z:
        names = set(z.namelist())
        raw = {n: z.read(n) for n in names if n in (
            'word/header1.xml', 'word/header2.xml',
            'word/_rels/header1.xml.rels', 'word/media/image1.jpeg')}
    if not all(k in raw for k in ('word/header1.xml', 'word/header2.xml')):
        return  # 模板若无页眉则跳过
    # 用独立文件名拷入，避免与底包已有部件重名（如 02 底自带空 header1.xml）
    blob = {
        'word/ht_header1.xml': raw['word/header1.xml'],
        'word/ht_header2.xml': raw['word/header2.xml'],
        'word/_rels/ht_header1.xml.rels': raw.get('word/_rels/header1.xml.rels', b''),
        'word/media/image1.jpeg': raw.get('word/media/image1.jpeg', b''),
    }
    blob = {k: v for k, v in blob.items() if v}
    with zipfile.ZipFile(str(merged_path)) as z:
        data = {n: z.read(n) for n in z.namelist()}
    doc_xml = ET.fromstring(data['word/document.xml'])
    rels_xml = ET.fromstring(data['word/_rels/document.xml.rels'])
    ct_xml = ET.fromstring(data['[Content_Types].xml'])

    def fresh_rid():
        nums = []
        for r in rels_xml:
            rid = r.get('Id', '')
            if rid.startswith('rId') and rid[3:].isdigit():
                nums.append(int(rid[3:]))
        fresh_rid.n = max(nums + [getattr(fresh_rid, 'n', 200)]) + 1
        return f'rId{fresh_rid.n}'

    # 1) 拷入部件（重名直接覆盖：ht_ 开头系本工具专用）
    for k, v in blob.items():
        data[k] = v
    empty_hdr = f'<w:hdr xmlns:w="{W_NS}"><w:p/></w:hdr>'.encode()
    data['word/header_empty.xml'] = empty_hdr
    # 2) rels
    RT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    rid_h1 = fresh_rid(); rid_h2 = fresh_rid(); rid_he = fresh_rid()
    for rid, tgt in [(rid_h1, 'ht_header1.xml'), (rid_h2, 'ht_header2.xml'),
                     (rid_he, 'header_empty.xml')]:
        ET.SubElement(rels_xml, f'{{{REL_NS}}}Relationship',
                      {'Id': rid, 'Type': RT + 'header', 'Target': tgt})
    # 3) Content_Types（含 jpeg兜底）
    have_parts = {e.get('PartName') for e in ct_xml}
    have_ext = {e.get('Extension') for e in ct_xml}
    for part in ['/word/ht_header1.xml', '/word/ht_header2.xml', '/word/header_empty.xml']:
        if part not in have_parts:
            ET.SubElement(ct_xml, f'{{{CT_NS}}}Override', {
                'PartName': part,
                'ContentType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml'})
    if 'jpeg' not in have_ext and 'jpg' not in have_ext:
        ET.SubElement(ct_xml, f'{{{CT_NS}}}Default',
                      {'Extension': 'jpeg', 'ContentType': 'image/jpeg'})
    # 4) 各节挂接（模板角色：header1=首页眉, header2=默认眉；无页脚）
    sects = [e for e in doc_xml.iter() if e.tag == f'{{{W_NS}}}sectPr']
    for i, sp in enumerate(sects):
        for old in [e for e in list(sp)
                    if e.tag in (f'{{{W_NS}}}headerReference', f'{{{W_NS}}}footerReference')]:
            sp.remove(old)
        if i == 0:
            refs = [('headerReference', 'first', rid_h1),
                    ('headerReference', 'default', rid_h2)]
        else:
            refs = [('headerReference', 'default', rid_he)]
        at = 0
        for tag, typ, rid in refs:
            el = ET.Element(f'{{{W_NS}}}{tag}',
                            {f'{{{W_NS}}}type': typ, f'{{{R_NS}}}id': rid})
            sp.insert(at, el)
            at += 1
        if i == 0 and sp.find(f'{{{W_NS}}}titlePg') is None:
            ET.SubElement(sp, f'{{{W_NS}}}titlePg')
    data['word/document.xml'] = ET.tostring(doc_xml, xml_declaration=True, encoding='UTF-8')
    data['word/_rels/document.xml.rels'] = ET.tostring(rels_xml, xml_declaration=True, encoding='UTF-8')
    data['[Content_Types].xml'] = ET.tostring(ct_xml, xml_declaration=True, encoding='UTF-8')
    with zipfile.ZipFile(str(merged_path), 'w', zipfile.ZIP_DEFLATED) as z:
        for n, v in data.items():
            z.writestr(n, v)


def _case_names(mapping: dict):
    safe_client = re.sub(r'[\\/:*?"<>|]', "", mapping["{{client_names}}"] or "委托人").replace("、", "_")[:30]
    safe_opp = re.sub(r'[\\/:*?"<>|]', "", mapping["{{opponent}}"] or "对方")[:20]
    base = f"{safe_client}与{safe_opp}{mapping['{{case_reason_full}}']}_{mapping['{{date_y}}']}-{mapping['{{date_m}}']}-{mapping['{{date_d}}']}"
    return safe_client, safe_opp, base


def _copy_section(src_doc, dst_doc):
    if src_doc.sections and dst_doc.sections:
        s0, s1 = src_doc.sections[0], dst_doc.sections[0]
        s1.page_width, s1.page_height = s0.page_width, s0.page_height
        s1.top_margin, s1.bottom_margin = s0.top_margin, s0.bottom_margin
        s1.left_margin, s1.right_margin = s0.left_margin, s0.right_margin


def generate_case(data: dict, cfg: dict) -> dict:
    """每案生成一个文件夹：6 份独立文档（可直接打印）+ 1 份合一版。
    返回 {"dir": Path, "merged": Path, "parts": [Path...]}。"""
    from docx import Document
    from docx.shared import Pt
    mapping = build_mapping(data, cfg)
    _, _, base = _case_names(mapping)
    case_dir = OUTPUT_DIR / base
    case_dir.mkdir(parents=True, exist_ok=True)

    filled = []  # (label, doc)
    parts = []
    for idx, (fname, label) in enumerate(ORDER, start=1):
        src_path = TEMPLATES_DIR / fname
        if not src_path.exists():
            continue
        doc = Document(str(src_path))
        replace_in_doc(doc, mapping)
        strip_doc(doc)
        _clear_footers(doc)  # 独立文件也不显示页码
        filled.append((label, doc))
        part_path = case_dir / f"{idx}_{label}.docx"
        doc.save(str(part_path))
        parts.append(part_path)

    # 合一版：无封面无标题，6 份正文各成一节、奇数页起始（原模板版式各自保留）。
    # 直接打印 = 6 份独立效果；双面打印每份都从正面开始。
    merged = _new_clean_doc()
    for i, (_, doc) in enumerate(filled):
        _append_doc(merged, doc)
        if i < len(filled) - 1:
            _section_break(merged, doc)
    if filled:
        _apply_section_to_body(merged, filled[-1][1])
    strip_doc(merged)
    # 文档以表格结尾时 Word 会隐式追加一个空段落并可能挤出空白尾页；
    # 主动加一个 1pt 零间距收尾段，把隐式段落“占住”，避免尾部空白页。
    from docx.shared import Pt as _Pt
    tail = _bare_paragraph(merged)
    tail.paragraph_format.space_before = _Pt(0)
    tail.paragraph_format.space_after = _Pt(0)
    tail.paragraph_format.line_spacing = 1.0
    _r = tail.add_run()
    _r.font.size = _Pt(1)
    merged_path = case_dir / f"0_合一版_6份齐全_{base}.docx"
    merged.save(str(merged_path))
    _clone_contract_hf(merged_path)  # 合同节挂回 logo 页眉；全文件无页码
    return {"dir": case_dir, "merged": merged_path, "parts": parts}


def generate_merged(data: dict, cfg: dict) -> Path:
    """兼容旧调用：返回合一版路径（同时也会生成6份独立文档）。"""
    return generate_case(data, cfg)["merged"]
